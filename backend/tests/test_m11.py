"""M11: demo ma'lumotlarni tozalash, takroriy importda tugagan shartnomalarni yopish, taklif xatlari, monitoring."""

import pytest
from django.core import mail
from django.core.management import call_command
from django.core.management.base import CommandError

from apps.accounts.models import Role, RolePermission, User
from apps.crm.models import Student
from apps.education.models import Contract
from apps.finance.models import Account, CashMovement, Payment
from apps.importer import reset, services
from apps.notifications.models import TelegramLink
from apps.references.models import AttendanceStatus, Branch, Course
from tests.test_finance import fin, pay  # noqa: F401 — fixture
from tests.test_importer import ON, xlsx  # noqa: F401 — fixture

pytestmark = pytest.mark.django_db


def test_prepare_production_keeps_only_essentials(api, fin, make_user):  # noqa: F811
    pay(api(fin["cashier"]), fin, amount=100_000)
    boss = fin["boss"]
    boss.set_password("eski-parol")
    boss.save()
    TelegramLink.objects.create(user=boss, chat_id=777)
    TelegramLink.objects.create(user=fin["cashier"], chat_id=778)
    roles, perms, statuses = Role.objects.count(), RolePermission.objects.count(), AttendanceStatus.objects.count()

    dry = reset.run([boss.email])  # sinov rejimi
    assert dry["To'lovlar"] == 1 and Payment.objects.count() == 1 and User.objects.count() > 1

    reset.run([boss.email], commit=True)
    assert list(User.objects.values_list("email", flat=True)) == [boss.email]
    boss.refresh_from_db()
    assert boss.check_password("eski-parol") and list(boss.branches.values_list("name", flat=True)) == ["Emirat"]
    assert TelegramLink.objects.get().chat_id == 777
    assert (Role.objects.count(), RolePermission.objects.count(), AttendanceStatus.objects.count()) == (roles, perms, statuses)
    assert list(Branch.objects.values_list("name", flat=True)) == ["Emirat"] and Branch.objects.get().default_cashbox is None
    for m in (Payment, CashMovement, Account, Contract, Student, Course):
        assert m._base_manager.count() == 0, m
    with pytest.raises(ValueError, match="topilmadi"):
        reset.run(["yoq@codial.test"])


def test_reimport_closes_contracts_ended_in_appsheet(xlsx, tmp_path):  # noqa: F811
    services.run(xlsx, on=ON, dry_run=False)
    assert Contract.objects.get(legacy_id="c2").end_date is None
    import datetime

    import openpyxl

    wb = openpyxl.load_workbook(xlsx)
    ws = wb["Contracts"]
    head = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2):
        if row[0].value == "c2":
            row[head.index("EndDate")].value = datetime.datetime(2026, 10, 20)
            row[head.index("StopReason")].value = "Moliyaviy qiyinchilik"
    p = tmp_path / "v2.xlsx"
    wb.save(p)
    rep = services.run(str(p), on=ON, dry_run=False)
    c2 = Contract.objects.get(legacy_id="c2")
    assert c2.end_date == datetime.date(2026, 10, 20) and c2.stop_reason == "Moliyaviy qiyinchilik"
    assert rep.counts["Shartnomalar (yopildi)"]["yangilandi"] == 1 and rep.checks["Tizimda faol shartnomalar (shu sanada)"] == 1


def test_send_invites_only_passwordless_active_staff(xlsx, make_user, settings):  # noqa: F811
    settings.EMAIL_BACKEND = "django.core.mail.backends.locmem.EmailBackend"
    services.run(xlsx, on=ON, dry_run=False)
    make_user("BOSH_MENEJER")  # paroli bor — xat ketmaydi
    call_command("send_invites")
    assert mail.outbox == []  # --send siz faqat ro'yxat
    call_command("send_invites", "--send")
    assert [m.to for m in mail.outbox] == [["ustoz1@mail.test"]]  # «Qo'shilmagan» yordamchiga ketmaydi
    assert "AppSheet o'rniga" in mail.outbox[0].body and "parol o'rnating" in mail.outbox[0].subject


def test_seed_demo_refuses_on_real_data(xlsx):  # noqa: F811
    services.run(xlsx, on=ON, dry_run=False)
    with pytest.raises(CommandError, match="haqiqiy"):
        call_command("seed_demo")


def test_health(client):
    assert client.get("/api/v1/health/").json() == {"ok": True}
