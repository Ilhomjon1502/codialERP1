"""M10: AppSheet importi — faqat faol shartnomalar va ularga kerakli ma'lumotlar, qarzsiz, takroriy ishga tushirishda dublikatsiz."""

import datetime

import openpyxl
import pytest

from apps.accounts.models import Role, User
from apps.crm.models import Student
from apps.education.models import Contract, Group
from apps.importer import services
from apps.references.models import Branch, Course, DayPattern

pytestmark = pytest.mark.django_db
ON = datetime.date(2026, 11, 1)
EM, XO = "301d8ce0", "75a97c52"

SHEETS = {
    "Branches": [("BranchID", "Name", "Status"), (EM, "Emirat", "Faol"), (XO, "Xonaqoh", "Nofaol")],
    "Courses": [("CourseID", "BranchID", "Name", "Status"), ("crs1", f"{EM}, {XO}", "Frontend", "Faol"), ("crs2", EM, "Android", "Nofaol")],
    "LessonDays": [("DayID", "Name", "WeekDay"), ("d1", "Monday , Wednesday , Friday", "1 , 3 , 5"), ("d2", "Saturday, Sunday", "6 , 0")],
    "LessonTimes": [("TimeID", "BranchID", "StartTime", "EndTime"), ("t1", EM, datetime.time(14, 0), datetime.time(16, 0))],
    "Rooms": [("RoomID", "BranchID", "Name", "RoomCapacity", "Status"), ("r1", EM, "1-xona", 12.0, "Faol")],
    "Users": [("UserID", "FullName", "RoleID", "BranchID", "Phone", "Team", "Email", "Status"),
              ("u1", "Ustoz Bir", "Ustoz", EM, None, "Ta'lim", "Ustoz1@Mail.test", "Qo'shilgan"),
              ("u2", "Yordamchi Ikki", "Yordmachi", EM, None, None, "yor2@mail.test", "Qo'shilmagan"),
              ("u3", "Boshqa Ustoz", "Ustoz", EM, None, None, "boss@mail.test", "Qo'shilgan")],
    "Groups": [("GroupID", "BranchID", "CourseID", "DayID", "TimeID", "TeacherID", "SupporterID", "RoomID", "StartDate", "EndDate", "GroupNumber", "Comment", "CreatedTime"),
               ("g1", EM, "crs1", "d1", "t1", "u1", "u2", "r1", datetime.datetime(2026, 9, 1), None, 7.0, None, datetime.datetime(2026, 8, 30, 10, 0)),
               ("g2", EM, "crs2", "d2", "t1", "u1", None, "r1", datetime.datetime(2025, 1, 1), datetime.datetime(2025, 6, 1), 3.0, None, None)],
    "Students": [("StudentID", "BranchID", "FirstName", "LastName", "ParentName", "Gender", "Phone1", "Phone2", "BirthDate", "ParentID", "CreatedTime"),
                 (1001.0, EM, "Ali", "Valiyev", "Vali o'g'li.", "Erkak", 901112233.0, 998901112244.0, "05.03.2010", "p1", None),
                 ("s2", EM, "Lola", "Karimova", None, "Ayol", 931234567.0, 931234567.0, None, None, None),
                 ("s3", EM, "Bekzod", "Tursunov", None, "Erkak", 12345.0, None, None, None, None),
                 ("s4", EM, "Eski", "Talaba", None, "Erkak", 901000000.0, None, None, None, None),
                 ("s5", XO, "Boshqa", "Filial", None, "Erkak", 901000001.0, None, None, None, None)],
    "Parents": [("ParentID", "BranchID", "FirstName", "SecondName", "Phone"), ("p1", EM, "Vali", "Valiyev", 905556677.0)],
    "Contracts": [("ContractID", "BranchID", "StudentID", "CourseID", "GroupID", "ContractType", "ContractNumber", "StartDate", "EndDate", "StopReason", "Comment_Reason", "CreatedTime"),
                  ("1001_1", EM, 1001.0, "crs1", "g1", "Pullik", 1.0, "01.09.2026", None, None, None, datetime.datetime(2026, 9, 1, 9, 0)),  # faol (matn sana)
                  ("c2", EM, "s2", "crs1", "g1", "Grant", 2.0, datetime.datetime(2026, 9, 5), None, "Kursni tamomladi", None, None),  # faol, sabab tashlanadi
                  ("c3", EM, "s3", "crs1", "g1", "Pullik", 3.0, datetime.datetime(2026, 9, 5), None, None, None, None),  # telefon yaroqsiz → xato
                  ("c4", EM, "s4", "crs2", "g2", "Pullik", 4.0, datetime.datetime(2025, 1, 1), datetime.datetime(2025, 6, 1), "Kursni tamomladi", None, None),  # nofaol
                  ("c5", EM, "s4", "crs1", "g1", "Pullik", 5.0, datetime.datetime(2026, 12, 1), None, None, None, None),  # rejada
                  ("c6", XO, "s5", "crs1", "g1", "Pullik", 6.0, datetime.datetime(2026, 9, 1), None, None, None, None)],  # boshqa filial
}


@pytest.fixture
def xlsx(tmp_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in SHEETS.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    p = tmp_path / "appsheet.xlsx"
    wb.save(p)
    return str(p)


def test_dry_run_writes_nothing(xlsx):
    rep = services.run(xlsx, on=ON, dry_run=True)
    assert rep.checks["Excel'da faol shartnomalar"] == 4 and rep.checks["Ko'chirilgan faol shartnomalar"] == 2
    assert Contract.objects.count() == 0 and Student.objects.count() == 0 and not User.objects.filter(legacy_id="u1").exists()


def test_import_only_active_contracts_and_dependencies(xlsx):
    rep = services.run(xlsx, on=ON, dry_run=False)
    assert set(Contract.objects.values_list("legacy_id", flat=True)) == {"1001_1", "c2", "c5"}  # rejada ham; nofaol, boshqa filial va xatoli — yo'q
    assert any("c3" in e and "telefon" in e for e in rep.errors)
    c1 = Contract.objects.get(legacy_id="1001_1")
    assert c1.start_date == datetime.date(2026, 9, 1) and c1.opening_balance == 0 and c1.number == 1 and c1.branch.name == "Emirat"
    assert Contract.objects.get(legacy_id="c2").type == "Grant" and Contract.objects.get(legacy_id="c2").stop_reason == ""
    s = c1.student
    assert (s.first_name, s.father_name, s.phone1, s.phone2, s.birth_date) == ("Ali", "Vali o'g'li", "901112233", "901112244", datetime.date(2010, 3, 5))
    assert s.parent.phone == "905556677" and Student.objects.get(legacy_id="s2").phone2 == ""  # bir xil telefon tashlandi
    g = Group.objects.get(legacy_id="g1")
    assert g.number == 7 and g.name == "Frontend / #7" and g.created_at.date() == datetime.date(2026, 8, 30)
    assert not Group.objects.filter(legacy_id="g2").exists() and not Course.objects.filter(name="Android").exists()  # faqat kerakligi
    assert g.day_pattern.weekdays == [1, 3, 5] and g.room.capacity == 12 and g.time.start_time == datetime.time(14, 0)
    t, sup = g.teacher, g.supporter
    assert t.email == "ustoz1@mail.test" and t.role.code == "USTOZ" and t.employment_status == "Faol" and not t.has_usable_password()
    assert sup.role.code == "YORDAMCHI" and sup.employment_status == "Ishdan ketgan"  # «Yordmachi» tanildi
    assert any("Qo'shilmagan" in w for w in rep.warnings) and any("to'xtatish sababi" in w for w in rep.warnings)
    assert not User.objects.filter(legacy_id="u3").exists()  # guruhga biriktirilmagan xodim ko'chirilmaydi
    assert Branch.objects.filter(name="Emirat").count() == 1


def test_rerun_is_idempotent_and_links_existing_user(xlsx, make_user):
    me = make_user("BOSH_MENEJER", email="ustoz1@mail.test", full_name="Mening akkauntim")  # tizimda allaqachon bor pochta
    services.run(xlsx, on=ON, dry_run=False)
    rep = services.run(xlsx, on=ON, dry_run=False)
    assert Contract.objects.count() == 3 and Group.objects.count() == 1 and Student.objects.count() == 3 and DayPattern.objects.count() == 1
    assert rep.counts["Shartnomalar"] == {"yangilandi": 3}
    me.refresh_from_db()
    assert me.legacy_id == "u1" and me.role.code == "BOSH_MENEJER" and me.full_name == "Mening akkauntim"  # roli va ismi o'zgarmadi


def test_reader_normalizes_values():
    from apps.importer import reader

    assert sorted({7 if int(x) == 0 else int(x) for x in reader.ids("6 , 0")}) == [6, 7]
    assert reader.phone(998901234567.0) == "901234567" and reader.phone("90 123-45-67") == "901234567" and reader.phone(123) == ""
    assert reader.date("05.03.2010") == datetime.date(2010, 3, 5) and reader.ident(1001.0) == "1001"
    assert Role.objects.filter(code="YORDAMCHI").exists()


def test_transition_extras_opening_leads_prices_kpi_accounts(tmp_path):
    """2026-10-18 qarori: boshlang'ich qarz (AppSheet formulasi), ochiq lidlar, narx va dars soni, KPI, g'azna qoldiqlari."""
    from apps.crm.models import Lead
    from apps.finance.models import Account
    from apps.finance.services import balances
    from apps.payroll.models import KPIRate
    from apps.references.models import CoursePrice, LessonFrequency

    extra = {
        "Students": [*SHEETS["Students"], ("s6", EM, "Yangi", "Lidov", None, "Erkak", 907776655.0, None, None, None, None)],
        "CoursePrices": [("PriceID", "BranchID", "CourseID", "Year", "Month", "Price", "ToExport"), ("pr1", EM, "crs1", 2026.0, 10.0, 1200000.0, f"{EM}crs1202610")],
        "LessonFrequency": [("FrequenceID", "BranchID", "DayID", "CourceID", "Year", "Month", "Frequency"), ("f1", EM, "d1", "crs1 , crsX", 2026.0, 10.0, 12.0)],
        "Attendance": [("AttendanceID", "BranchID", "CourseID", "GroupID", "Date", "ContractID", "Status", "Year", "Month", "DayID"),
                       *[(f"a{i}", EM, "crs1", "g1", None, "1001_1", st, 2026.0, 10.0, "d1") for i, st in enumerate(["1", "X", "0", "-"])],
                       ("a9", EM, "crs1", "g1", None, "c2", "1", 2026.0, 10.0, "d1")],  # Grant — qarz yo'q
        "Payments": [("PaymentID", "BranchID", "ContractID", "Summa", "Discount", "Type", "Cashbox"), ("p1", EM, "1001_1", 350000.0, 50000.0, "Cash", "acc1")],
        "LeadSources": [("SourceID", "Name"), ("src1", "Instagram"), ("src2", "")],
        "Leads": [("LeadID", "StudentID", "Tags", "Noutbook", "AgeGroup", "StudentSocialStatus", "SourceID", "BranchID", "CourseID", "DayID", "TimeID",
                   "Status", "GroupID", "DemoDate", "Demo_Come", "CreatedTime", "CreatedBy"),
                  ("l1", "s6", "grand, tez", "Bor", "13-15", "Maktab o'quvchisi", "src1", EM, "crs1", "d1", "t1", "3.Sinov darsga yozildi", "g1", "20.10.2026", "", None, None),
                  ("l2", 1001.0, "", "Bor", "13-15", "", "src1", EM, "crs1", "d1", "t1", "5.O'qishni boshladi", "g1", None, "True", None, None)],
        "LeadComments": [("CommentID", "BranchID", "LeadID", "Comment", "CreatedTime", "CreatedBy"), ("lc1", EM, "l1", "Qo'ng'iroq qilindi", None, None)],
        "TeacherKPI": [("KPIID", "BranchID", "Status", "TeacherSupporter", "CourseID", "Year", "Month", "Type", "Fixed", "Percent"),
                       ("k1", EM, "Ustoz", "u1", "crs1", 2026.0, 10.0, "O'zgarmas", 120000.0, 0.0)],
        "Accounts": [("AccountID", "Name", "Responsible", "InitialBalance", "Status", "Cashtype"),
                     ("acc1", "Emirat g'azna", None, 100000.0, "Faol", "Cash"), ("acc2", "Bosh g'azna", None, None, "Faol", "All"), ("acc3", "Eski", None, 5.0, "Nofaol", "All")],
        "Transactions": [("TransactionID", "SenderCashbox", "Sum", "PaymentType", "RecieverCashbox", "ConfirmationStatus"),
                         ("t1", "acc1", 50000.0, "Cash", "acc2", "Tasdiqlandi"), ("t2", "acc1", 70000.0, "Cash", "acc2", "Tasdiqlashga")],
        "Expenses": [("ExpenseID", "Summa", "TypeID", "Cashbox"), ("e1", 30000.0, "Card", "acc2")],
        "SalaryGiving": [("SalaryID", "Summa", "CashType", "CashBox")], "CashFlow": [("CashID", "Summa", "CashType", "Cashbox")],
        "Exchange": [("ExchangeID", "Cashbox", "CashTypeFrom", "Summa", "CashTypeTo", "SummaTo"), ("x1", "acc2", "Cash", 20000.0, "Card", 19000.0)],
    }
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in {**SHEETS, **extra}.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    path = str(tmp_path / "full.xlsx")
    wb.save(path)
    emirat, bosh = Account.objects.create(name="Emirat"), Account.objects.create(name="Bosh gʻazna", cash_type="ALL")

    rep = services.run(path, on=datetime.date(2026, 10, 19), dry_run=False)
    assert [e for e in rep.errors if "c3" not in e] == []  # c3 — ataylab yaroqsiz telefon
    # 3 dars × 1 200 000 / 12 = 300 000; to'lov 350 000 + chegirma 50 000 → −100 000 (oldindan to'lagan); Grant — 0
    assert Contract.objects.get(legacy_id="1001_1").opening_balance == -100000 and Contract.objects.get(legacy_id="c2").opening_balance == 0
    lead = Lead.objects.get(legacy_id="l1")
    assert lead.status == "3" and lead.group.legacy_id == "g1" and lead.tags == ["grand", "tez"] and lead.source.name == "Instagram"
    assert lead.comments.get().comment == "Qo'ng'iroq qilindi" and not Lead.objects.filter(legacy_id="l2").exists()  # 5-status ko'chmaydi
    assert CoursePrice.objects.get().price == 1200000
    f = LessonFrequency.objects.get(legacy_id="f1")
    assert f.status == "Tasdiqlangan" and f.frequency == 12 and [c.name for c in f.courses.all()] == ["Frontend"]
    assert KPIRate.objects.get(legacy_id="k1").fixed_amount == 120000
    b = balances()
    assert b[emirat.pk]["balance"] == 100000 + 350000 - 50000 and b[emirat.pk]["by_type"] == {"CASH": 400000}
    assert b[bosh.pk]["by_type"] == {"CASH": 50000 - 20000, "CARD": -30000 + 19000}
    services.run(path, on=datetime.date(2026, 10, 19), dry_run=False)  # takroriy — dublikatsiz
    assert Lead.objects.count() == 1 and balances()[emirat.pk]["balance"] == 400000 and KPIRate.objects.count() == 1
