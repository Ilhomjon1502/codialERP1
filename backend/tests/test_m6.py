"""M6: ikki yo'nalishli baholash, ota-onalar bilan aloqa, Excel eksport; olib tashlangan bo'limlar (CLAUDE.md 7.4, 7.6)."""

import datetime
import io

import pytest
from openpyxl import load_workbook

from apps.accounts.models import RolePermission
from apps.core.models import AuditLog
from apps.core.roles import SECTION_KEYS
from apps.education.models import ParentCall
from tests.test_education import TODAY, world  # noqa: F401 — fixture

pytestmark = pytest.mark.django_db


def test_removed_sections():
    for gone in ("tasks", "discount50", "grading"):
        assert gone not in SECTION_KEYS
    assert not RolePermission.objects.filter(section__in=["tasks", "discount50", "grading"]).exists()


def test_teacher_evaluation(api, make_user, world):  # noqa: F811
    adm = api(make_user("ADMINISTRATOR"))
    c = world["contracts"][0]
    body = {"contract": str(c.pk), "s1": 9, "s2": 10, "s3": 8, "s4": 7, "s5": 9, "comment": "Yaxshi"}
    r = adm.post("/api/v1/evaluations/", body, format="json")
    assert r.status_code == 201, r.content
    assert r.json()["teacher"] == str(world["teacher"].pk) and r.json()["total"] == 43 and r.json()["scores"] == [9, 10, 8, 7, 9]
    assert adm.post("/api/v1/evaluations/", {**body, "s1": 11}, format="json").status_code == 400
    assert api(world["teacher"]).post("/api/v1/evaluations/", body, format="json").status_code == 403  # ustoz o'zini baholamaydi
    assert len(api(make_user("BOSH_MENEJER")).get(f"/api/v1/evaluations/?teacher={world['teacher'].pk}").json()) == 1


def test_parent_calls(api, make_user, world):  # noqa: F811
    adm = api(make_user("ADMINISTRATOR"))
    c = world["contracts"][0]
    r = adm.post("/api/v1/parent-calls/", {"contract": str(c.pk), "comment": "Darsga 2 marta kelmadi"}, format="json")
    assert r.status_code == 201, r.content
    d = r.json()
    assert d["group"] == str(world["group"].pk) and d["course"] == str(world["course"].pk) and d["phone1"] == c.student.phone1 and d["edited"] is False
    assert adm.post("/api/v1/parent-calls/", {"contract": str(c.pk), "comment": " "}, format="json").status_code == 400
    future = (TODAY + datetime.timedelta(days=1)).isoformat()
    assert adm.post("/api/v1/parent-calls/", {"contract": str(c.pk), "comment": "x", "date": future}, format="json").status_code == 400
    ParentCall.objects.filter(pk=d["id"]).update(created_at=d["created_at"])
    e = adm.patch(f"/api/v1/parent-calls/{d['id']}/", {"comment": "Ota-onasi bilan gaplashildi", "contract": str(world["contracts"][1].pk)}, format="json")
    assert e.status_code == 200 and e.json()["comment"] == "Ota-onasi bilan gaplashildi" and e.json()["contract"] == str(c.pk)  # talaba o'zgarmaydi
    assert adm.delete(f"/api/v1/parent-calls/{d['id']}/").status_code == 403  # o'chirish — faqat Admin va Bosh menejer
    assert api(world["teacher"]).get("/api/v1/parent-calls/").json()[0]["id"] == d["id"]
    assert api(make_user("USTOZ")).get("/api/v1/parent-calls/").json() == []
    assert api(make_user("BOSH_MENEJER")).delete(f"/api/v1/parent-calls/{d['id']}/").status_code == 204


def test_excel_export(api, make_user, world):  # noqa: F811
    bm = api(make_user("BOSH_MENEJER"))
    r = bm.get(f"/api/v1/export/contracts/?groups={world['group'].pk}")
    assert r.status_code == 200 and r["Content-Type"].startswith("application/vnd.openxmlformats")
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws["A1"].value == "№" and ws.max_row == 4 and ws["B2"].value.startswith("Talaba")
    assert AuditLog.objects.filter(action="export").exists()
    assert api(make_user("YORDAMCHI")).get("/api/v1/export/contracts/").status_code == 403
    assert api(make_user("USTOZ")).get("/api/v1/export/contracts/").status_code == 200  # ustoz — faqat o'z guruhlari (bo'sh)
