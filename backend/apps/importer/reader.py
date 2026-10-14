"""AppSheet (Google Sheets) .xlsx eksportini o'qish: varaqlar → lug'atlar ro'yxati, ID/sana/telefon/son normallashtirish."""

import datetime
import re

import openpyxl


class Workbook:
    def __init__(self, path, sheets):
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        missing = [s for s in sheets if s not in wb.sheetnames]
        if missing:
            raise ValueError(f"Faylda varaq yo'q: {', '.join(missing)}")
        self.rows = {}
        for name in sheets:
            it = wb[name].iter_rows(values_only=True)
            head = [str(h).strip() if h is not None else "" for h in next(it, ())]
            self.rows[name] = [dict(zip(head, r, strict=False)) for r in it if any(v not in (None, "") for v in r)]
        wb.close()

    def __getitem__(self, name):
        return self.rows[name]

    def by_id(self, name, key):
        return {ident(r.get(key)): r for r in self.rows[name] if ident(r.get(key))}


def ident(v):
    """AppSheet kaliti: '81a9d571', 1001.0 → '1001'."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def ids(v):
    """Ro'yxat ustuni: '301d8ce0 , 75a97c52' → ['301d8ce0', '75a97c52']."""
    return [x for x in (p.strip() for p in str(v or "").split(",")) if x]


def text(v):
    return re.sub(r"\s+", " ", str(v)).strip() if v not in (None,) else ""


def integer(v, default=None):
    if v in (None, ""):
        return default
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return default


DATE_FORMATS = ("%d.%m.%Y", "%m/%d/%Y", "%Y-%m-%d", "%d.%m.%Y %H:%M:%S", "%m/%d/%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S")


def date(v):
    if v in (None, ""):
        return None
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    t = str(v).strip()
    for f in DATE_FORMATS:
        try:
            return datetime.datetime.strptime(t, f).date()
        except ValueError:
            pass
    raise ValueError(f"sana tushunarsiz: {t[:20]}")


def moment(v):
    if isinstance(v, datetime.datetime):
        return v
    d = date(v)
    return datetime.datetime.combine(d, datetime.time()) if d else None


def time(v):
    if isinstance(v, datetime.time):
        return v.replace(second=0, microsecond=0)
    if isinstance(v, datetime.datetime):
        return v.time().replace(second=0, microsecond=0)
    m = re.match(r"^\s*(\d{1,2}):(\d{2})", str(v or ""))
    if not m:
        raise ValueError(f"vaqt tushunarsiz: {str(v)[:10]}")
    return datetime.time(int(m.group(1)), int(m.group(2)))


def phone(v):
    """9 raqam (+998 siz). 998901234567 → 901234567. Yaroqsiz bo'lsa ''."""
    if v in (None, ""):
        return ""
    s = str(int(v)) if isinstance(v, float) else str(v)
    digits = re.sub(r"\D", "", s)
    if len(digits) == 12 and digits.startswith("998"):
        digits = digits[3:]
    return digits if len(digits) == 9 else ""
