"""Shartnomalar ro'yxatini .xlsx ga yuklash (CLAUDE.md 7.6, AppSheet'dagi «Shartnoma yuklash» boti o'rniga)."""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import services

HEAD = ["№", "Talaba", "1-telefon", "2-telefon", "Vasiy", "Vasiy telefoni", "Kurs", "Guruh", "Turi", "Boshlanish", "Tugash", "To'xtatish sababi",
        "Status", "Hisoblangan", "To'langan", "Chegirma", "Balans"]
MONEY_COLS = {14, 15, 16, 17}


def contracts_xlsx(contracts):
    """contracts — with_balance() qilingan, student/parent/group yuklangan queryset."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Shartnomalar"
    ws.append(HEAD)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="F26A1B")
        cell.alignment = Alignment(vertical="center")
    for c in contracts:
        st, p = c.student, c.student.parent
        ws.append([c.number, st.full_name, st.phone1, st.phone2, str(p) if p else "", p.phone if p else "", str(c.course), str(c.group), c.type,
                   c.start_date, c.end_date, c.stop_reason, services.period_status(c.start_date, c.end_date),
                   int(c.charged), int(c.paid), int(c.discount_sum), int(c.balance)])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if cell.column in MONEY_COLS:
                cell.number_format = "#,##0"
            elif cell.column in (10, 11):
                cell.number_format = "DD.MM.YYYY"
    widths = [6, 32, 12, 12, 22, 14, 18, 22, 9, 12, 12, 30, 9, 13, 13, 12, 13]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
